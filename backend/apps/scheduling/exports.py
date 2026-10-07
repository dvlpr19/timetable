"""Excel and PDF export of a group / teacher / room timetable in a chosen language."""

import io
from dataclasses import dataclass

from django.template.loader import render_to_string
from django.utils import timezone, translation
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.academics.models import Academy, BlockedPeriod, LessonTime
from apps.core.dates import format_long_date, weekday_name

from .models import Schedule, ScheduleEntry


@dataclass
class Table:
    title: str
    subtitle: str
    academy: str
    columns: list[str]  # header row (first column = row label)
    rows: list[list[str]]  # cells: newline-separated lines
    closed: set[tuple[int, int]]  # (row, column) of blocked cells, e.g. Friday prayer


def cell_text(entry: ScheduleEntry, view: str) -> str:
    a = entry.assignment
    lines = [f"{a.subject.name} ({a.lesson_type.name.lower()})"]
    place = entry.room.name if entry.room else _("Online")
    who = {
        "group": a.teacher.short_name,
        "teacher": a.target_name,
        "room": f"{a.target_name}, {a.teacher.short_name}",
    }[view]
    lines.append(f"{place} · {who}")
    if entry.week_parity in ("odd", "even"):
        lines.append(_("Odd weeks") if entry.week_parity == "odd" else _("Even weeks"))
    return "\n".join(lines)


def build_table(schedule: Schedule, entries: list[ScheduleEntry], view: str, name: str) -> Table:
    titles = {
        "group": _("Timetable of group %(name)s"),
        "teacher": _("Timetable of %(name)s"),
        "room": _("Timetable of room %(name)s"),
    }
    semester = schedule.semester
    subtitle = f"{semester.year.name}, {semester.get_kind_display().lower()}"
    academy = Academy.objects.first()
    dated = any(e.date for e in entries)
    hard_blocks = list(BlockedPeriod.objects.filter(is_hard=True))
    closed: set[tuple[int, int]] = set()

    def blocked_name(weekday: int, number: int) -> str:
        lt = next((t for t in times if t.number == number), None)
        for b in hard_blocks:
            if lt and b.applies_to(weekday, lt.form_id) and b.overlaps(lt):
                return b.name
        return ""

    form_ids = {e.lesson_time.form_id for e in entries}
    times = list(LessonTime.objects.filter(form_id__in=form_ids).order_by("number", "start"))
    numbers = sorted({lt.number for lt in times})
    span = {}
    for lt in times:
        span.setdefault(lt.number, f"{lt.start:%H:%M}–{lt.end:%H:%M}")

    if dated:
        days = sorted({e.date for e in entries})
        columns = [_("Date"), *[f"{n}\n{span[n]}" for n in numbers]]
        rows = []
        for day in days:
            row = [format_long_date(day)]
            for n in numbers:
                text = "\n\n".join(
                    cell_text(e, view)
                    for e in entries
                    if e.date == day and e.lesson_time.number == n
                )
                if not text and (name_ := blocked_name(day.weekday(), n)):
                    text = name_
                    closed.add((len(rows), len(row)))
                row.append(text)
            rows.append(row)
    else:
        weekdays = sorted({e.weekday for e in entries} | set(range(6)))
        columns = [_("Lesson"), *[weekday_name(d) for d in weekdays]]
        rows = []
        for n in numbers:
            row = [f"{n}\n{span[n]}"]
            for d in weekdays:
                text = "\n\n".join(
                    cell_text(e, view)
                    for e in entries
                    if e.weekday == d and e.lesson_time.number == n
                )
                if not text and (name_ := blocked_name(d, n)):
                    text = name_
                    closed.add((len(rows), len(row)))
                row.append(text)
            rows.append(row)
    return Table(
        title=titles[view] % {"name": name},
        subtitle=subtitle,
        academy=academy.name if academy else "",
        columns=columns,
        rows=rows,
        closed=closed,
    )


def to_xlsx(table: Table) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = table.title[:31]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    width = len(table.columns)
    ws.cell(row=1, column=1, value=table.academy).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=table.title).font = Font(bold=True, size=14)
    ws.cell(row=3, column=1, value=table.subtitle).font = Font(color="5B6863")
    ws.cell(row=1, column=width, value=signature_text()).alignment = Alignment(
        horizontal="right", wrap_text=True, vertical="top"
    )
    thin = Side(style="thin", color="E1E5DF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="0B5D4B")
    closed_fill = PatternFill("solid", fgColor="EEF1EC")
    for c, title in enumerate(table.columns, start=1):
        cell = ws.cell(row=5, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 16 if c == 1 else 30
    for r, row in enumerate(table.rows, start=6):
        lines = max((v.count("\n") + 1 for v in row), default=1)
        ws.row_dimensions[r].height = max(30, 15 * lines)
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value or None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if c == 1:
                cell.font = Font(bold=True)
            if (r - 6, c - 1) in table.closed:
                cell.fill = closed_fill
                cell.font = Font(italic=True, color="5B6863")
    ws.cell(row=7 + len(table.rows), column=1, value=generated_text()).font = Font(
        color="5B6863", size=9
    )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(table: Table) -> bytes:
    from weasyprint import HTML

    html = render_to_string(
        "scheduling/timetable_pdf.html",
        {
            "t": table,
            "rows": [
                [{"text": v, "closed": (r, c) in table.closed} for c, v in enumerate(row)]
                for r, row in enumerate(table.rows)
            ],
            "signature": signature_text(),
            "signature_line": _("(signature)"),
            "generated": generated_text(),
            "lang": translation.get_language(),
        },
    )
    return HTML(string=html).write_pdf()


def signature_text() -> str:
    return _("Approved by: vice-rector for academic affairs")


def generated_text() -> str:
    return _("Generated on %(date)s") % {"date": format_long_date(timezone.localdate())}


def export(schedule, entries, view, name, fmt, lang) -> tuple[bytes, str]:
    with translation.override(lang):
        table = build_table(schedule, list(entries), view, name)
        if fmt == "pdf":
            return to_pdf(table), "application/pdf"
        return to_xlsx(table), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


__all__ = ["Table", "build_table", "export", "to_pdf", "to_xlsx"]
