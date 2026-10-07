"""Reports: numbers agree with the timetable, files open, only staff may see them."""

import io

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.academics.models import Faculty, Room
from apps.accounts.models import User
from apps.scheduling.models import EntryOccurrence, Schedule

pytestmark = pytest.mark.django_db


def client_for(username: str) -> APIClient:
    client = APIClient()
    client.force_authenticate(User.objects.get(username=username))
    return client


def test_plan_report_counts_real_dates(demo):
    data = client_for("dekanat").get("/api/reports/plan/").data
    published = Schedule.objects.get(status="published")
    held = EntryOccurrence.objects.filter(schedule=published, status="scheduled").count()
    assert data["totals"]["scheduled"] == held
    row = next(r for r in data["rows"] if r["scheduled"])
    assert row["difference"] == row["scheduled"] - row["planned"]
    assert {c["key"] for c in data["columns"]} >= {"planned", "scheduled", "lost", "percent"}


def test_faculty_filter_keeps_only_its_groups(demo):
    rus = Faculty.objects.get(code="RUS")
    rows = client_for("admin").get(f"/api/reports/plan/?faculty={rus.pk}").data["rows"]
    assert rows and all("-R" in r["target"] or "R-" in r["target"] for r in rows)


def test_teacher_report_hours_are_two_per_lesson(demo):
    rows = client_for("kafedra").get("/api/reports/teachers/").data["rows"]
    busy = next(r for r in rows if r["scheduled_lessons"])
    assert busy["hours"] == 2 * busy["scheduled_lessons"]
    assert busy["norm_hours"] > 0 and busy["weekly"] > 0


def test_room_report_uses_weekly_slots_without_closed_times(demo):
    data = client_for("admin").get("/api/reports/rooms/").data
    assert len(data["rows"]) == Room.objects.filter(is_active=True).count()
    used = next(r for r in data["rows"] if r["weekly"])
    # kunduzgi: 6 days x 6 lessons - 2 closed by the Friday prayer; kechki: 5 days x 2
    assert used["available"] == 6 * 6 - 2 + 5 * 2
    assert used["percent"] == round(100 * used["weekly"] / used["available"])
    assert 0 < used["fill"] <= 100


def test_files_carry_title_semester_and_signature(demo):
    c = client_for("admin")
    xlsx = c.get("/api/reports/teachers/?type=xlsx&lang=ru")
    assert xlsx.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws["B2"].value is None and ws["A2"].value  # title
    assert "2025-2026" in ws["A3"].value
    assert ws.cell(row=1, column=ws.max_column).value  # signature place
    pdf = c.get("/api/reports/rooms/?type=pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_reports_are_for_staff_only(demo):
    assert client_for("talaba").get("/api/reports/plan/").status_code == 403
    assert client_for("oqituvchi").get("/api/reports/rooms/").status_code == 403
    assert client_for("admin").get("/api/reports/nonsense/").status_code == 400
    assert client_for("admin").get("/api/reports/plan/?type=doc").status_code == 400
