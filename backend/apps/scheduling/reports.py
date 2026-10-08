"""Reports for the academic office: plan fulfilment, teacher workload, room occupancy.

Each report is a plain table (columns + rows + totals) shown on screen as JSON and exported
to Excel / PDF with the academy name, semester and a place for the signature.
"""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field

from django.db.models import Q
from django.template.loader import render_to_string
from django.utils import translation
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.academics.choices import WeekParity
from apps.academics.models import Academy, Group, Room, Teacher, TeachingAssignment

from .exports import generated_text, signature_text
from .models import EntryOccurrence, OccurrenceStatus, Schedule
from .validators import load_context, placements_from_entries, plan_fulfilment

# One lesson (para) is 80 minutes = 2 academic hours (ASSUMPTIONS.md).
HOURS_PER_LESSON = 2


@dataclass
class Column:
    key: str
    label: str
    numeric: bool = False


@dataclass
class Report:
    kind: str
    title: str
    subtitle: str
    academy: str
    columns: list[Column]
    rows: list[dict]
    totals: dict | None = None
    notes: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "kind": self.kind,
            "title": self.title,
            "subtitle": self.subtitle,
            "academy": self.academy,
            "columns": [c.__dict__ for c in self.columns],
            "rows": self.rows,
            "totals": self.totals,
            "notes": self.notes,
        }


def _subtitle(schedule: Schedule, faculty) -> str:
    semester = schedule.semester
    parts = [f"{semester.year.name}, {semester.get_kind_display().lower()}"]
    if schedule.status != "published":  # say which draft the numbers are about
        parts.append(schedule.name)
    if faculty:
        parts.append(faculty.name)
    return " · ".join(parts)


def _academy() -> str:
    academy = Academy.objects.first()
    return academy.name if academy else ""


def _pct(part: float, whole: float) -> float | None:
    return round(100 * part / whole) if whole else None


def _faculty_groups(faculty) -> set[int] | None:
    if faculty is None:
        return None
    return set(
        Group.objects.filter(program_form__program__faculty=faculty).values_list("id", flat=True)
    )


# --------------------------------------------------------------------------- plan fulfilment


def plan_report(schedule: Schedule, faculty=None) -> Report:
    """Planned lessons of each assignment vs lessons that really take place (dates)."""
    ctx = load_context(schedule.semester)
    placements = placements_from_entries(schedule.entries.all())
    groups = _faculty_groups(faculty)
    rows_by_id = {r.assignment_id: r for r in plan_fulfilment(placements, ctx)}
    assignments = (
        TeachingAssignment.objects.filter(pk__in=rows_by_id)
        .select_related(
            "subject",
            "lesson_type",
            "teacher",
            "group",
            "subgroup__group",
            "stream",
            "period__form",
        )
        .prefetch_related("stream__groups")
        .order_by("period__form__order", "subject__name_uz", "lesson_type__code")
    )
    rows, totals = [], defaultdict(int)
    for a in assignments:
        info = ctx.assignments[a.pk]
        if groups is not None and not set(info.group_ids) & groups:
            continue
        r = rows_by_id[a.pk]
        diff = r.scheduled - r.planned
        rows.append(
            {
                "subject": a.subject.name,
                "lesson_type": a.lesson_type.name,
                "target": a.target_name,
                "teacher": a.teacher.short_name,
                "form": a.period.form.name,
                "planned": r.planned,
                "scheduled": r.scheduled,
                "lost": r.lost,
                "difference": diff,
                "percent": _pct(r.scheduled, r.planned),
                "status": "ok" if diff == 0 else ("less" if diff < 0 else "more"),
            }
        )
        for k in ("planned", "scheduled", "lost"):
            totals[k] += rows[-1][k]
    totals["difference"] = totals["scheduled"] - totals["planned"]
    totals["percent"] = _pct(totals["scheduled"], totals["planned"])
    matching = sum(1 for r in rows if r["status"] == "ok")
    return Report(
        kind="plan",
        title=_("Curriculum fulfilment"),
        subtitle=_subtitle(schedule, faculty),
        academy=_academy(),
        columns=[
            Column("subject", _("Subject")),
            Column("lesson_type", _("Lesson type")),
            Column("target", _("Groups")),
            Column("teacher", _("Teacher")),
            Column("form", _("Form")),
            Column("planned", _("Planned"), True),
            Column("scheduled", _("Scheduled"), True),
            Column("lost", _("Lost (days off, cancelled)"), True),
            Column("difference", _("Difference"), True),
            Column("percent", _("Fulfilment") + ", %", True),
        ],
        rows=rows,
        totals={"subject": _("Total"), **totals},
        notes=[
            _("Lessons are counted on real dates: holidays and cancelled lessons are lost."),
            _("%(ok)s of %(all)s assignments match the plan exactly.")
            % {"ok": matching, "all": len(rows)},
        ],
    )


# --------------------------------------------------------------------------- teacher workload


def teachers_report(schedule: Schedule, faculty=None) -> Report:
    """Weekly lessons against the limit and semester hours against the norm."""
    teachers = Teacher.objects.select_related("department").order_by("last_name", "first_name")
    if faculty:
        teachers = teachers.filter(department__faculty=faculty)
    weekly: dict[int, float] = defaultdict(float)
    for teacher_id, parity in schedule.entries.filter(date__isnull=True).values_list(
        "assignment__teacher_id", "week_parity"
    ):
        weekly[teacher_id] += 1 if parity == WeekParity.EVERY else 0.5
    planned: dict[int, int] = defaultdict(int)
    for teacher_id, total in TeachingAssignment.objects.filter(
        period__semester=schedule.semester
    ).values_list("teacher_id", "total_lessons"):
        planned[teacher_id] += total
    held: dict[int, int] = defaultdict(int)
    for teacher_id in EntryOccurrence.objects.filter(
        schedule=schedule, status=OccurrenceStatus.SCHEDULED
    ).values_list("teacher_id", flat=True):
        held[teacher_id] += 1
    rows, totals = [], defaultdict(float)
    for t in teachers:
        if not (planned[t.pk] or held[t.pk] or weekly[t.pk]):
            continue
        norm = t.annual_load_hours / 2  # one semester
        hours = held[t.pk] * HOURS_PER_LESSON
        rows.append(
            {
                "teacher": t.full_name,
                "department": t.department.name,
                "position": t.get_position_display(),
                "weekly": weekly[t.pk],
                "weekly_max": t.max_weekly_lessons,
                "planned_lessons": planned[t.pk],
                "scheduled_lessons": held[t.pk],
                "hours": hours,
                "norm_hours": norm,
                "percent": _pct(hours, norm),
                "status": "over" if weekly[t.pk] > t.max_weekly_lessons else "ok",
            }
        )
        for k in ("weekly", "planned_lessons", "scheduled_lessons", "hours", "norm_hours"):
            totals[k] += rows[-1][k]
    totals["percent"] = _pct(totals["hours"], totals["norm_hours"])
    over = sum(1 for r in rows if r["status"] == "over")
    return Report(
        kind="teachers",
        title=_("Teacher workload"),
        subtitle=_subtitle(schedule, faculty),
        academy=_academy(),
        columns=[
            Column("teacher", _("Teacher")),
            Column("department", _("Department")),
            Column("position", _("Position")),
            Column("weekly", _("Lessons a week"), True),
            Column("weekly_max", _("Weekly limit"), True),
            Column("planned_lessons", _("Planned lessons"), True),
            Column("scheduled_lessons", _("Scheduled lessons"), True),
            Column("hours", _("Hours"), True),
            Column("norm_hours", _("Semester norm, hours"), True),
            Column("percent", _("Of the norm") + ", %", True),
        ],
        rows=rows,
        totals={"teacher": _("Total"), **totals},
        notes=[
            _("One lesson = %(n)s academic hours; the semester norm is half the annual load.")
            % {"n": HOURS_PER_LESSON},
            _("Odd/even-week lessons count as half a lesson a week."),
            _("Teachers over their weekly limit: %(n)s.") % {"n": over},
        ],
    )


# --------------------------------------------------------------------------- room occupancy


def rooms_report(schedule: Schedule, faculty=None) -> Report:
    """How much of the week each room is used and how full it is."""
    ctx = load_context(schedule.semester)
    # weekly slots a room can host: every (weekday, lesson time) of in-person weekly forms,
    # minus closed times such as the Friday prayer
    slots = set()
    for form in ctx.forms.values():
        if not form.requires_room or form.mode != "weekly":
            continue
        for lt in ctx.lesson_times.values():
            if lt.form_id != form.id:
                continue
            for wd in form.study_weekdays:
                if not any(b.is_hard and b.hits(wd, form.id, lt) for b in ctx.blocked):
                    slots.add((wd, lt.start, lt.end))
    available = len(slots)

    rooms = Room.objects.filter(is_active=True).select_related("building", "room_type")
    if faculty:
        rooms = rooms.filter(Q(faculty=faculty) | Q(faculty__isnull=True))
    used: dict[int, float] = defaultdict(float)
    fill: dict[int, list[float]] = defaultdict(list)
    session: dict[int, int] = defaultdict(int)
    for p in placements_from_entries(schedule.entries.filter(room__isnull=False)):
        room = ctx.rooms[p.room_id]
        students = ctx.assignments[p.assignment_id].student_count
        if p.is_dated:
            session[p.room_id] += 1
            continue
        used[p.room_id] += 1 if p.week_parity == WeekParity.EVERY else 0.5
        fill[p.room_id].append(students / room.capacity)
    rows, totals = [], defaultdict(float)
    for r in rooms.order_by("building__name", "floor", "name"):
        avg_fill = round(100 * sum(fill[r.pk]) / len(fill[r.pk])) if fill[r.pk] else None
        rows.append(
            {
                "room": r.name,
                "building": r.building.name,
                "room_type": r.room_type.name,
                "capacity": r.capacity,
                "weekly": used[r.pk],
                "available": available,
                "percent": _pct(used[r.pk], available),
                "fill": avg_fill,
                "session": session[r.pk],
            }
        )
        totals["weekly"] += used[r.pk]
        totals["available"] += available
        totals["session"] += session[r.pk]
    totals["percent"] = _pct(totals["weekly"], totals["available"])
    idle = sum(1 for r in rows if not r["weekly"] and not r["session"])
    return Report(
        kind="rooms",
        title=_("Room occupancy"),
        subtitle=_subtitle(schedule, faculty),
        academy=_academy(),
        columns=[
            Column("room", _("Room")),
            Column("building", _("Building")),
            Column("room_type", _("Room type")),
            Column("capacity", _("Seats"), True),
            Column("weekly", _("Lessons a week"), True),
            Column("available", _("Weekly slots"), True),
            Column("percent", _("Occupancy") + ", %", True),
            Column("fill", _("Average fill") + ", %", True),
            Column("session", _("Session lessons"), True),
        ],
        rows=rows,
        totals={"room": _("Total"), **totals},
        notes=[
            _(
                "Weekly slots: study days × lesson times of in-person weekly forms, without "
                "closed times (Friday prayer)."
            ),
            _("Average fill: students of a lesson divided by the room's seats."),
            _("Rooms without any lesson: %(n)s.") % {"n": idle},
        ],
    )


BUILDERS = {"plan": plan_report, "teachers": teachers_report, "rooms": rooms_report}


# --------------------------------------------------------------------------- files


def _fmt(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def report_to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    width = len(report.columns)
    ws.cell(row=1, column=1, value=report.academy).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=report.title).font = Font(bold=True, size=14)
    ws.cell(row=3, column=1, value=report.subtitle).font = Font(color="5B6478")
    ws.cell(row=1, column=width, value=signature_text()).alignment = Alignment(
        horizontal="right", wrap_text=True, vertical="top"
    )
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="1565C0")
    for c, col in enumerate(report.columns, start=1):
        cell = ws.cell(row=5, column=c, value=col.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 12 if col.numeric else 26
    ws.row_dimensions[5].height = 32
    body = [*report.rows, *([report.totals] if report.totals else [])]
    for r, row in enumerate(body, start=6):
        is_total = report.totals is not None and r == 5 + len(body)
        for c, col in enumerate(report.columns, start=1):
            value = row.get(col.key)
            cell = ws.cell(row=r, column=c, value=value if col.numeric else (value or ""))
            cell.border = border
            cell.alignment = Alignment(
                horizontal="right" if col.numeric else "left", vertical="top", wrap_text=True
            )
            if is_total:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="EAF2FC")
    line = 7 + len(body)
    for note in report.notes:
        ws.cell(row=line, column=1, value=note).font = Font(color="5B6478", size=9)
        line += 1
    ws.cell(row=line, column=1, value=generated_text()).font = Font(color="5B6478", size=9)
    ws.freeze_panes = "B6"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def report_to_pdf(report: Report) -> bytes:
    from weasyprint import HTML

    def cells(row):
        return [{"text": _fmt(row.get(c.key)), "numeric": c.numeric} for c in report.columns]

    html = render_to_string(
        "scheduling/report_pdf.html",
        {
            "r": report,
            "rows": [cells(row) for row in report.rows],
            "totals": cells(report.totals) if report.totals else None,
            "signature": signature_text(),
            "signature_line": _("(signature)"),
            "generated": generated_text(),
            "lang": translation.get_language(),
        },
    )
    return HTML(string=html).write_pdf()


def build(kind: str, schedule: Schedule, faculty, lang: str) -> Report:
    with translation.override(lang):
        return BUILDERS[kind](schedule, faculty)


def render_file(kind: str, schedule: Schedule, faculty, fmt: str, lang: str) -> tuple[bytes, str]:
    with translation.override(lang):
        report = BUILDERS[kind](schedule, faculty)
        if fmt == "pdf":
            return report_to_pdf(report), "application/pdf"
        return (
            report_to_xlsx(report),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
