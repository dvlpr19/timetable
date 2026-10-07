import io

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.academics.models import Room, Teacher
from apps.accounts.models import User

pytestmark = pytest.mark.django_db


def client(username="admin"):
    c = APIClient()
    c.force_authenticate(User.objects.get(username=username))
    return c


def filled_template(c, url, rows, lang="uz"):
    response = c.get(url, HTTP_ACCEPT_LANGUAGE=lang)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    out.name = "data.xlsx"
    return out


def test_room_template_has_titles_in_user_language(demo):
    response = client().get("/api/rooms/import-template/", HTTP_ACCEPT_LANGUAGE="ru")
    ws = load_workbook(io.BytesIO(response.content)).active
    assert ws["A1"].value == "Корпус"
    assert ws["A2"].value == "building"
    assert ws.row_dimensions[2].hidden


def test_import_creates_then_updates_rooms(demo):
    c = client()
    rows = [
        ("A", "A-401", 4, "Auditoriya", 32, "ha", 0, "ha"),
        ("B bino", "B-410", 4, "Компьютерный класс", 20, "yes", 20, "1"),
    ]
    upload = filled_template(c, "/api/rooms/import-template/", rows)
    r = c.post("/api/rooms/import/", {"file": upload}, format="multipart")
    assert r.status_code == 200, r.data
    assert r.data == {"created": 2, "updated": 0}
    assert Room.objects.get(name="B-410").room_type.code == "computer"

    upload = filled_template(
        c, "/api/rooms/import-template/", [("A", "A-401", 4, "auditorium", 40)]
    )
    r = c.post("/api/rooms/import/", {"file": upload}, format="multipart")
    assert r.data == {"created": 0, "updated": 1}
    assert Room.objects.get(name="A-401").capacity == 40


def test_one_bad_row_saves_nothing(demo):
    c = client()
    rows = [
        ("A", "A-501", 5, "auditorium", 30),
        ("Z", "Z-1", 1, "auditorium", 30),  # no building Z
        ("A", "A-502", 5, "auditorium", "ko'p"),  # not a number
    ]
    upload = filled_template(c, "/api/rooms/import-template/", rows)
    r = c.post(
        "/api/rooms/import/", {"file": upload}, format="multipart", HTTP_ACCEPT_LANGUAGE="uz"
    )
    assert r.status_code == 400
    assert [(e["row"], e["column"]) for e in r.data["errors"]] == [(4, "building"), (5, "capacity")]
    assert "topilmadi" in r.data["errors"][0]["message"]
    assert not Room.objects.filter(name__in=["A-501", "A-502"]).exists()


def test_kafedra_cannot_import_teachers_of_other_departments(demo):
    c = client("kafedra")
    rows = [
        ("Yangiyev", "Olim", "", "QH", "teacher", "none", "staff", 900, 16, "uz,ar", "TAF"),
        ("Begona", "Ali", "", "IGT", "teacher", "none", "staff", 900, 16, "uz", "ARB"),
    ]
    upload = filled_template(c, "/api/teachers/import-template/", rows)
    r = c.post("/api/teachers/import/", {"file": upload}, format="multipart")
    assert r.status_code == 400
    assert [e["row"] for e in r.data["errors"]] == [4]
    assert not Teacher.objects.filter(last_name__in=["Yangiyev", "Begona"]).exists()


def test_not_an_excel_file(demo):
    upload = io.BytesIO(b"not excel")
    upload.name = "x.xlsx"
    r = client().post("/api/rooms/import/", {"file": upload}, format="multipart")
    assert r.status_code == 400
