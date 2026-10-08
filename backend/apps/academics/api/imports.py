"""Excel import for rooms, subjects, teachers, groups and students.

Template: row 1 = column titles in the user's language, row 2 = hidden column keys, data
from row 3. Files without the key row are accepted too if row 1 holds titles in any of
the three languages. Related objects are matched by code or by name in any language.
The whole file is imported in one transaction: one bad row and nothing is saved.
"""

import io
from collections.abc import Callable
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.db import models, transaction
from django.http import HttpResponse
from django.utils import translation
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import status
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from ..models import (
    Building,
    Department,
    EducationForm,
    Faculty,
    Group,
    Program,
    ProgramForm,
    Room,
    RoomType,
    Student,
    SubGroup,
    Subject,
    Teacher,
)
from ..naming import group_name

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TRUE = {"1", "true", "yes", "ha", "да", "+", "x"}
FALSE = {"0", "false", "no", "yo'q", "yoq", "нет", "-", ""}


class RowError(Exception):
    def __init__(self, column: str, message: str):
        super().__init__(message)
        self.column = column
        self.message = message


def find_by_code_or_name(model, value: str, *, code_field: str = "code"):
    value = str(value).strip()
    q = models.Q(**{f"{code_field}__iexact": value})
    if hasattr(model, "name_uz"):
        for lang in ("uz", "ru", "en"):
            q |= models.Q(**{f"name_{lang}__iexact": value})
    found = list(model.objects.filter(q)[:2])
    if len(found) != 1:
        raise ValueError(_("“%(value)s” was not found.") % {"value": value})
    return found[0]


def to_bool(value) -> bool:
    text = str(value if value is not None else "").strip().lower()
    if text in TRUE:
        return True
    if text in FALSE:
        return False
    raise ValueError(_("Write yes or no."))


def to_list(value) -> list[str]:
    return [p.strip() for p in str(value or "").replace(";", ",").split(",") if p.strip()]


@dataclass
class Column:
    key: str
    field: str  # model field whose verbose_name titles the column
    parse: Callable | None = None  # value -> python value (raise ValueError for bad input)
    required: bool = True


@dataclass
class Resource:
    model: type[models.Model]
    columns: list[Column]
    key: tuple[str, ...]  # natural key for update-or-create
    build: Callable[[dict], dict] = lambda row: row  # parsed row -> model kwargs
    after_save: Callable | None = None
    m2m: dict = field(default_factory=dict)


def _groups_build(row: dict) -> dict:
    faculty = row.pop("faculty")
    program = Program.objects.filter(faculty=faculty, code=str(row.pop("program"))).first()
    if program is None:
        raise RowError("program", _("This program does not exist in the faculty."))
    pf = ProgramForm.objects.filter(program=program, form=row.pop("form")).first()
    if pf is None:
        raise RowError("form", _("The program is not offered in this education form."))
    row["program_form"] = pf
    if not row.get("name"):
        row["name"] = group_name(pf.group_prefix, row["course"], row["number"])
    return row


def _groups_after(group: Group, row: dict) -> None:
    if not group.subgroups.exists():
        half = (group.student_count + 1) // 2
        SubGroup.objects.create(group=group, number=1, student_count=half)
        SubGroup.objects.create(group=group, number=2, student_count=group.student_count - half)


def _students_build(row: dict) -> dict:
    number = row.pop("subgroup", None)
    row["subgroup"] = (
        SubGroup.objects.filter(group=row["group"], number=int(number)).first() if number else None
    )
    return row


RESOURCES: dict[str, Resource] = {
    "rooms": Resource(
        Room,
        [
            Column("building", "building", lambda v: find_by_code_or_name(Building, v)),
            Column("name", "name"),
            Column("floor", "floor", int),
            Column("room_type", "room_type", lambda v: find_by_code_or_name(RoomType, v)),
            Column("capacity", "capacity", int),
            Column("has_projector", "has_projector", to_bool, required=False),
            Column("computer_count", "computer_count", int, required=False),
            Column("is_active", "is_active", to_bool, required=False),
        ],
        key=("building", "name"),
    ),
    "subjects": Resource(
        Subject,
        [
            Column("code", "code"),
            Column("name_uz", "name_uz"),
            Column("name_ru", "name_ru", required=False),
            Column("name_en", "name_en", required=False),
            Column("department", "department", lambda v: find_by_code_or_name(Department, v)),
            Column("is_language", "is_language", to_bool, required=False),
        ],
        key=("code",),
    ),
    "teachers": Resource(
        Teacher,
        [
            Column("last_name", "last_name"),
            Column("first_name", "first_name"),
            Column("middle_name", "middle_name", required=False),
            Column("department", "department", lambda v: find_by_code_or_name(Department, v)),
            Column("position", "position"),
            Column("degree", "degree", required=False),
            Column("employment", "employment", required=False),
            Column("annual_load_hours", "annual_load_hours", int),
            Column("max_weekly_lessons", "max_weekly_lessons", int, required=False),
            Column("teaching_languages", "teaching_languages", to_list),
            Column(
                "subjects",
                "subjects",
                lambda v: [find_by_code_or_name(Subject, c) for c in to_list(v)],
                required=False,
            ),
        ],
        key=("last_name", "first_name", "middle_name"),
        m2m={"subjects": "subjects"},
    ),
    "groups": Resource(
        Group,
        [
            Column("faculty", "program_form", lambda v: find_by_code_or_name(Faculty, v)),
            Column("program", "program_form"),
            Column("form", "program_form", lambda v: find_by_code_or_name(EducationForm, v)),
            Column("course", "course", int),
            Column("number", "number", int),
            Column("name", "name", required=False),
            Column("teaching_language", "teaching_language"),
            Column("student_count", "student_count", int),
            Column("gender_composition", "gender_composition", required=False),
            Column("shift", "shift", int, required=False),
        ],
        key=("name",),
        build=_groups_build,
        after_save=_groups_after,
    ),
    "students": Resource(
        Student,
        [
            Column("hemis_id", "hemis_id"),
            Column("last_name", "last_name"),
            Column("first_name", "first_name"),
            Column("middle_name", "middle_name", required=False),
            Column(
                "group",
                "group",
                lambda v: find_by_code_or_name(Group, v, code_field="name"),
            ),
            Column("subgroup", "subgroup", required=False),
            Column("gender", "gender"),
            Column("phone", "phone", required=False),
        ],
        key=("hemis_id",),
        build=_students_build,
    ),
}

COLUMN_TITLE_OVERRIDES = {"faculty": "faculty", "program": "program", "form": "education form"}


def column_title(resource: Resource, column: Column) -> str:
    if column.key in COLUMN_TITLE_OVERRIDES and column.field == "program_form":
        return _(COLUMN_TITLE_OVERRIDES[column.key]).capitalize()
    return str(resource.model._meta.get_field(column.field).verbose_name).capitalize()


def build_template(resource: Resource) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = str(resource.model._meta.verbose_name_plural)[:31]
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1565C0")
    for i, col in enumerate(resource.columns, start=1):
        title = column_title(resource, col) + ("" if col.required else f" ({_('optional')})")
        cell = ws.cell(row=1, column=i, value=title)
        cell.font, cell.fill = bold, fill
        ws.cell(row=2, column=i, value=col.key)
        ws.column_dimensions[cell.column_letter].width = max(14, len(title) + 2)
    ws.row_dimensions[2].hidden = True
    ws.freeze_panes = "A3"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _header_map(resource: Resource, rows: list[tuple]) -> tuple[dict[int, Column], int]:
    """Column index → Column, and the first data row index."""
    keys = {c.key: c for c in resource.columns}
    if len(rows) > 1 and any(str(v or "").strip() in keys for v in rows[1]):
        return {
            i: keys[str(v).strip()] for i, v in enumerate(rows[1]) if str(v or "").strip() in keys
        }, 2
    titles: dict[str, Column] = {}
    for lang in ("uz", "ru", "en"):
        with translation.override(lang):
            for c in resource.columns:
                titles[column_title(resource, c).lower()] = c
    header = {}
    for i, v in enumerate(rows[0] if rows else ()):
        title = str(v or "").split(" (")[0].strip().lower()
        if title in titles:
            header[i] = titles[title]
    return header, 1


def import_rows(resource: Resource, data: bytes, in_scope: Callable) -> tuple[dict, list[dict]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    header, start = _header_map(resource, rows)
    missing = [c.key for c in resource.columns if c.required and c not in header.values()]
    if missing:
        titles = ", ".join(column_title(resource, c) for c in resource.columns if c.key in missing)
        return {}, [
            {"row": 1, "column": "", "message": _("Missing columns: %(c)s") % {"c": titles}}
        ]

    errors: list[dict] = []
    counts = {"created": 0, "updated": 0}
    with transaction.atomic():
        for offset, raw in enumerate(rows[start:], start=start + 1):
            if not any(v not in (None, "") for v in raw):
                continue
            try:
                obj, created = _import_row(resource, header, raw)
                if not in_scope(obj):
                    raise RowError("", _("You may not change this record."))
                counts["created" if created else "updated"] += 1
            except RowError as e:
                errors.append({"row": offset, "column": e.column, "message": e.message})
        if errors:
            transaction.set_rollback(True)
    return counts, errors


def _import_row(resource: Resource, header: dict[int, Column], raw: tuple):
    parsed: dict = {}
    for i, col in header.items():
        value = raw[i] if i < len(raw) else None
        if value in (None, ""):
            if col.required:
                raise RowError(col.key, _("This field is required."))
            continue
        try:
            parsed[col.key] = col.parse(value) if col.parse else str(value).strip()
        except (ValueError, TypeError) as e:
            raise RowError(col.key, str(e) or _("Wrong value.")) from e
    m2m = {name: parsed.pop(key) for key, name in resource.m2m.items() if key in parsed}
    kwargs = resource.build(parsed)
    lookup = {k: kwargs.get(k, "") for k in resource.key}
    obj = resource.model.objects.filter(**lookup).first()
    created = obj is None
    obj = obj or resource.model()
    for name, value in kwargs.items():
        setattr(obj, name, value)
    try:
        obj.full_clean(validate_unique=created)
        with transaction.atomic():
            obj.save()
    except ValidationError as e:
        column, messages = next(iter(e.message_dict.items()))
        raise RowError(column, " ".join(messages)) from e
    except Exception as e:  # database constraint
        raise RowError("", str(e).splitlines()[0]) from e
    for name, value in m2m.items():
        getattr(obj, name).set(value)
    if resource.after_save:
        resource.after_save(obj, kwargs)
    return obj, created


class ExcelImportMixin:
    """Adds GET …/import-template/ and POST …/import/ (multipart, field "file")."""

    import_resource: str = ""

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        resource = RESOURCES[self.import_resource]
        response = HttpResponse(build_template(resource), content_type=XLSX)
        response["Content-Disposition"] = f'attachment; filename="{self.import_resource}.xlsx"'
        return response

    @action(detail=False, methods=["post"], url_path="import", parser_classes=[MultiPartParser])
    def import_file(self, request):
        upload = request.FILES.get("file")
        if upload is None:
            return Response({"detail": _("Choose an Excel file.")}, status=400)
        try:
            counts, errors = import_rows(
                RESOURCES[self.import_resource], upload.read(), self.in_scope
            )
        except Exception:  # not an xlsx file
            return Response({"detail": _("The file could not be read as Excel (.xlsx).")}, 400)
        if errors:
            return Response({"errors": errors}, status=status.HTTP_400_BAD_REQUEST)
        return Response(counts)
